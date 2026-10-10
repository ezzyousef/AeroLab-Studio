"""Importing awkward files, and the Excel and Origin writers."""
from __future__ import annotations

import json
import math

import numpy as np
import pytest

from aerolab.core import measurements as M, validation as V
from aerolab.io import excel as X, origin as O, readers as R


# ------------------------------------------------------------------- delimited text
def test_plain_csv(tmp_path):
    path = tmp_path / "plain.csv"
    path.write_text("Strain,Stress\n0,0\n1,2\n2,4\n", encoding="utf-8")
    ds = R.read_any(path)[0]
    assert ds.headers == ["Strain", "Stress"]
    assert ds.n_rows == 3
    assert ds.columns[1][2] == pytest.approx(4.0)


def test_instrument_preamble_is_skipped(tmp_path):
    path = tmp_path / "instron.csv"
    path.write_text(
        "# Instron 5943\n# Gauge length: 20 mm\n\n"
        "Time (s),Tensile strain (%),Tensile stress (MPa)\n"
        "0,0,0\n1,1.5,0.3\n2,3.0,0.6\n", encoding="utf-8")
    ds = R.read_any(path)[0]
    assert ds.headers == ["Time", "Tensile strain", "Tensile stress"]
    assert ds.units == ["s", "%", "MPa"]
    assert len(ds.meta["preamble"]) == 2
    assert ds.n_rows == 3


def test_units_row_and_european_decimals(tmp_path):
    path = tmp_path / "origin.txt"
    path.write_text("Temperature\tWeight\nC\t%\n30,000\t99,99\n40,000\t99,50\n",
                    encoding="utf-8")
    ds = R.read_any(path)[0]
    assert ds.units == ["C", "%"]
    assert ds.meta["decimal"] == ","
    assert ds.columns[0][0] == pytest.approx(30.0)
    assert ds.columns[1][1] == pytest.approx(99.50)


def test_semicolon_separator(tmp_path):
    path = tmp_path / "semi.csv"
    path.write_text("A;B\n1;2\n3;4\n", encoding="utf-8")
    ds = R.read_any(path)[0]
    assert ds.meta["delimiter"] == ";"
    assert ds.n_cols == 2


def test_whitespace_separated(tmp_path):
    path = tmp_path / "space.txt"
    path.write_text("X   Y\n1.0   2.0\n3.0   4.0\n", encoding="utf-8")
    ds = R.read_any(path)[0]
    assert ds.n_cols == 2 and ds.n_rows == 2


def test_nested_bracket_units_are_split(tmp_path):
    path = tmp_path / "nested.csv"
    path.write_text("Thermal conductivity (mW/(m.K)),Density (g/cm^3)\n31.2,0.157\n",
                    encoding="utf-8")
    ds = R.read_any(path)[0]
    assert ds.headers == ["Thermal conductivity", "Density"]
    assert ds.units == ["mW/(m.K)", "g/cm^3"]


def test_text_column_is_kept_as_labels(tmp_path):
    path = tmp_path / "named.csv"
    path.write_text("Sample,Density\nNFT-1,0.157\nNFT-2,0.149\n", encoding="utf-8")
    ds = R.read_any(path)[0]
    assert ds.text_columns["Sample"] == ["NFT-1", "NFT-2"]
    assert ds.columns[1][0] == pytest.approx(0.157)


def test_missing_values_become_nan(tmp_path):
    path = tmp_path / "gaps.csv"
    path.write_text("A,B\n1,2\n3,\n5,n/a\n", encoding="utf-8")
    ds = R.read_any(path)[0]
    assert math.isnan(ds.columns[1][1]) and math.isnan(ds.columns[1][2])


def test_missing_file_raises():
    with pytest.raises(R.ImportError_):
        R.read_any("no_such_file_anywhere.csv")


def test_empty_file_raises(tmp_path):
    path = tmp_path / "empty.csv"
    path.write_text("", encoding="utf-8")
    with pytest.raises(R.ImportError_):
        R.read_any(path)


# ----------------------------------------------------------------------- json, npz
def test_json_records(tmp_path):
    path = tmp_path / "records.json"
    path.write_text(json.dumps([{"x": 1, "y": 2}, {"x": 3, "y": 4}]), encoding="utf-8")
    ds = R.read_any(path)[0]
    assert ds.headers == ["x", "y"] and ds.n_rows == 2


def test_json_columns(tmp_path):
    path = tmp_path / "cols.json"
    path.write_text(json.dumps({"Strain (%)": [0, 1], "Stress (MPa)": [0, 2]}),
                    encoding="utf-8")
    ds = R.read_any(path)[0]
    assert ds.headers == ["Strain", "Stress"] and ds.units == ["%", "MPa"]


def test_npz_round_trip(tmp_path):
    path = tmp_path / "arrays.npz"
    np.savez(path, **{"Strain (%)": np.arange(5.0), "Stress (MPa)": np.arange(5.0) * 2})
    ds = R.read_any(path)[0]
    assert set(ds.headers) == {"Strain", "Stress"} and ds.n_rows == 5


# --------------------------------------------------------------------------- excel
def test_excel_round_trip(tmp_path):
    import openpyxl
    frame = [("Strain (%)", "Stress (MPa)"), (0, 0), (1, 2), (2, 4)]
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Run 1"
    for row in frame:
        sheet.append(list(row))
    path = tmp_path / "book.xlsx"
    book.save(path)

    datasets = R.read_any(path)
    assert len(datasets) == 1
    assert datasets[0].headers == ["Strain", "Stress"]
    assert datasets[0].meta["sheet"] == "Run 1"


def test_every_demo_dataset_imports_and_analyses(samples_dir):
    for measurement in M.MEASUREMENTS:
        datasets = R.read_any(samples_dir / measurement.sample)
        assert datasets, measurement.sample
        dataset = datasets[0]
        mapping = measurement.match_columns(dataset.headers)
        missing = [c.key for c in measurement.required_channels()
                   if mapping.get(c.key) is None]
        assert not missing, f"{measurement.id}: could not match {missing} in {dataset.headers}"


# -------------------------------------------------------------------- excel export
def test_workbook_has_every_sheet_and_chart(tmp_path, analysed):
    import openpyxl

    report = X.ExcelReport(project="test", author="tester", validation=V.run_all(),
                           manual=[{"name": "Porosity", "inputs": "a=1", "value": 80.0,
                                    "unit": "%", "reference": "paper"}])
    for measurement in M.MEASUREMENTS:
        result = analysed[measurement.id]
        if isinstance(result, list):
            for i, one in enumerate(result):
                report.add(X.RunRecord(f"row {i + 1}", measurement, one))
        else:
            report.add(X.RunRecord(measurement.sample, measurement, result))

    path = X.write_workbook(report, tmp_path / "out.xlsx")
    assert path.exists() and path.stat().st_size > 50_000

    book = openpyxl.load_workbook(path)
    for required in ("Summary", "Calculations", "Validation", "Charts", "Provenance"):
        assert required in book.sheetnames
    assert len(book["Summary"]["A"]) > 100, "the summary should list every metric"
    assert len(getattr(book["Charts"], "_charts", [])) >= 10

    statuses = [book["Validation"].cell(r, 5).value
                for r in range(5, book["Validation"].max_row + 1)]
    assert "FAILED" not in statuses


def test_table_measurements_share_one_sheet(tmp_path, analysed):
    import openpyxl
    measurement = M.get("thermal_point")
    report = X.ExcelReport(include_charts=False)
    for i, result in enumerate(analysed["thermal_point"]):
        report.add(X.RunRecord(f"sample {i + 1}", measurement, result))
    path = X.write_workbook(report, tmp_path / "table.xlsx")
    book = openpyxl.load_workbook(path)
    assert measurement.name in book.sheetnames
    sheet = book[measurement.name]
    assert sheet.max_row - 5 == len(analysed["thermal_point"])


def test_sheet_names_stay_within_excels_limit(tmp_path, analysed):
    import openpyxl
    measurement = M.get("stress_strain")
    report = X.ExcelReport(include_charts=False)
    for i in range(3):
        report.add(X.RunRecord("a-very-long-sample-name-that-will-not-fit" + str(i),
                               measurement, analysed["stress_strain"]))
    path = X.write_workbook(report, tmp_path / "names.xlsx")
    for name in openpyxl.load_workbook(path).sheetnames:
        assert len(name) <= 31


# ------------------------------------------------------------------- origin export
def test_origin_availability_reports_a_reason():
    ok, why = O.origin_available()
    assert isinstance(ok, bool) and why


def test_labtalk_package_is_complete(tmp_path, analysed):
    items = [O.OriginExportItem(m.sample, m, analysed[m.id])
             for m in M.MEASUREMENTS if m.plots]
    result = O.write_script_package(items, tmp_path / "pkg")
    assert result.mode == "script"
    assert result.books == len(items)

    script = (tmp_path / "pkg" / "AeroLab_build.ogs").read_text(encoding="utf-8")
    for expected in ("newbook", "impasc", "plotxy", "range rr", "label -xb",
                     "rescale", "legendupdate", "save -i"):
        assert expected in script, f"the script is missing {expected!r}"

    # Each of these was a real defect found by running the script inside Origin.
    assert "options.names." not in script, \
        "explicit impasc name options make Origin import zero rows"
    assert "[Main]" not in script, \
        "a [Main] section stops Origin executing the file when it is opened"
    assert "page.longname$" in script, \
        "a bare impasc overwrites the book long name with the file path"
    assert "set %C " not in script, \
        "set %C only reaches the first plot in a layer; each plot needs its own range"
    for line in script.splitlines():
        if line.startswith("plotxy"):
            assert "[%(" in line, f"ranges must use the captured short name: {line}"

    csvs = sorted((tmp_path / "pkg" / "data").glob("*.csv"))
    assert len(csvs) == len(items)
    assert (tmp_path / "pkg" / "README.txt").exists()

    head = csvs[0].read_text(encoding="utf-8").splitlines()
    assert len(head) > 4, "a CSV should have three header rows and data"


def test_origin_csv_carries_long_name_units_and_comments(tmp_path, analysed):
    measurement = M.get("crossover")
    item = O.OriginExportItem("sweep", measurement, analysed["crossover"])
    O.write_script_package([item], tmp_path / "pkg")
    lines = (tmp_path / "pkg" / "data" / f"{item.safe_id}.csv") \
        .read_text(encoding="utf-8").splitlines()
    assert lines[0].split(",")[0] == "Oscillation stress"      # long name
    assert lines[1].split(",")[0] == "Pa"                       # units
    assert "G" in lines[2]                                      # comments: the series label
    float(lines[3].split(",")[0])                               # then numbers


def test_export_falls_back_to_a_script_when_asked(tmp_path, analysed):
    measurement = M.get("stress_strain")
    item = O.OriginExportItem("s", measurement, analysed["stress_strain"])
    result = O.export_to_origin([item], tmp_path / "out", force_script=True)
    assert result.mode == "script"
    assert (tmp_path / "out" / "AeroLab_build.ogs").exists()


def test_export_with_nothing_to_do_is_harmless(tmp_path):
    result = O.export_to_origin([], tmp_path / "empty")
    assert result.books == 0 and result.messages


def test_uncommented_instrument_preamble_is_skipped(tmp_path):
    path = tmp_path / "instron.csv"
    path.write_text("Sample: NFT-4\nOperator: lab\nDate: 2024-03-01\n\n"
                    "Strain (%),Stress (MPa)\n0,0\n1,0.2\n2,0.4\n3,0.5\n", encoding="utf-8")
    ds = R.read_any(path)[0]
    assert ds.headers == ["Strain", "Stress"] and ds.units == ["%", "MPa"]
    assert ds.n_rows == 4
    assert any("NFT-4" in line for line in ds.meta["preamble"])


def test_whitespace_header_keeps_units_with_their_names(tmp_path):
    path = tmp_path / "tga.dat"
    path.write_text("Temperature (°C)  Weight (%)\n30 100\n100 98.5\n200 97\n", encoding="utf-8")
    ds = R.read_any(path)[0]
    assert ds.headers == ["Temperature", "Weight"] and ds.units == ["°C", "%"]


def test_number_parsing_does_not_invent_values():
    assert R._to_float("1.5e-3", ",") == pytest.approx(1.5e-3)
    assert R._to_float("1.234,5", ",") == pytest.approx(1234.5)
    assert R._to_float("12,345.6", ".") == pytest.approx(12345.6)
    assert math.isnan(R._to_float("1,5", "."))          # ambiguous: not silently 15
    assert R._to_float("1.0D+03", ".") == pytest.approx(1000.0)
    assert math.isnan(R._to_float("done", "."))


def test_old_xls_gives_a_clear_message(tmp_path):
    path = tmp_path / "old.xls"
    path.write_bytes(b"\xd0\xcf\x11\xe0not really")
    with pytest.raises(R.ImportError_, match="xlsx"):
        R.read_any(path)


def test_workbook_with_a_chartsheet_still_imports(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    import openpyxl.chart
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Strain (%)", "Stress (MPa)"])
    for i in range(5):
        ws.append([i * 1.0, i * 0.1])
    chart = openpyxl.chart.ScatterChart()
    chart.add_data(openpyxl.chart.Reference(ws, min_col=2, min_row=1, max_row=6), titles_from_data=True)
    wb.create_chartsheet("Chart").add_chart(chart)
    path = tmp_path / "with_chart.xlsx"
    wb.save(path)
    datasets = R.read_any(path)
    assert len(datasets) == 1 and datasets[0].n_rows == 5


def test_long_curves_are_charted_whole_and_ids_are_unique(tmp_path, analysed):
    import numpy as np
    from aerolab.core.curves import AnalysisResult
    m = M.get("stress_strain")
    x = np.linspace(0, 100, 5000)
    res = AnalysisResult("stress_strain", [], {"stress-strain": (x, x * 0.01)}, {})
    report = X.ExcelReport(project="t", include_curves=True, include_charts=True)
    report.add(X.RunRecord("A", m, res))
    X.write_workbook(report, tmp_path / "t.xlsx")
    openpyxl = pytest.importorskip("openpyxl")
    book = openpyxl.load_workbook(tmp_path / "t.xlsx", read_only=True)
    headers = [str(c) for ws in book.worksheets for row in ws.iter_rows(max_row=8, values_only=True)
               for c in row if c]
    assert any("thinned for chart" in h for h in headers)
    items = [O.OriginExportItem("same", m, res), O.OriginExportItem("same", m, res)]
    O._assign_ids(items)
    assert items[0].safe_id != items[1].safe_id


def test_nan_in_an_early_row_does_not_hide_the_header(tmp_path):
    path = tmp_path / "gap.csv"
    path.write_text("Strain (%),Stress (MPa)\n0,0\n1,NaN\n2,2\n3,3\n4,4\n", encoding="utf-8")
    ds = R.read_any(path)[0]
    assert ds.headers == ["Strain", "Stress"]
    assert ds.n_rows == 5
